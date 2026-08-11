from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ──────────────────────────────────────────────
section = doc.sections[0]
section.page_width  = Inches(8.27)
section.page_height = Inches(11.69)
section.left_margin   = Inches(1.0)
section.right_margin  = Inches(1.0)
section.top_margin    = Inches(1.0)
section.bottom_margin = Inches(1.0)

# ── Colour palette ────────────────────────────────────────────
NAVY     = RGBColor(0x1a, 0x3a, 0x5c)
NAVYDARK = RGBColor(0x0f, 0x25, 0x40)
GOLD     = RGBColor(0xc9, 0x96, 0x2a)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF3, 0xF5, 0xF8)
TEXT     = RGBColor(0x1e, 0x2d, 0x40)
SUBTEXT  = RGBColor(0x6a, 0x78, 0x88)
GREEN    = RGBColor(0x1a, 0x60, 0x30)
RED      = RGBColor(0xc0, 0x39, 0x2b)

# ── Helper: set cell background colour ────────────────────────
def set_cell_bg(cell, rgb: RGBColor):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    hex_color = '{:02X}{:02X}{:02X}'.format(rgb[0], rgb[1], rgb[2])
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def set_cell_border(cell, top=None, bottom=None, left=None, right=None):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side, val in [('top', top),('bottom', bottom),('left', left),('right', right)]:
        if val:
            el = OxmlElement('w:'+side)
            el.set(qn('w:val'),   val.get('val','single'))
            el.set(qn('w:sz'),    val.get('sz','6'))
            el.set(qn('w:space'),'0')
            el.set(qn('w:color'), val.get('color','auto'))
            tcBorders.append(el)
    tcPr.append(tcBorders)

def paragraph_shade(para, rgb: RGBColor):
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    hex_color = '{:02X}{:02X}{:02X}'.format(rgb[0], rgb[1], rgb[2])
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    pPr.append(shd)

# ── Helper: add a styled paragraph ────────────────────────────
def add_para(doc, text='', bold=False, italic=False, size=11, color=None,
             align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=6,
             font_name='Calibri'):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        run.font.name = font_name
        if color:
            run.font.color.rgb = color
    return p

def add_run(para, text, bold=False, italic=False, size=11, color=None, font_name='Calibri'):
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = font_name
    if color:
        run.font.color.rgb = color
    return run

# ══════════════════════════════════════════════════════════════
#  COVER BLOCK  (simulated with a table)
# ══════════════════════════════════════════════════════════════
cover = doc.add_table(rows=1, cols=1)
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
cover.style = 'Table Grid'
cell = cover.cell(0, 0)
set_cell_bg(cell, NAVYDARK)
set_cell_border(cell,
    top    ={'val':'single','sz':'18','color':'C9962A'},
    bottom ={'val':'single','sz':'18','color':'C9962A'},
    left   ={'val':'single','sz':'18','color':'C9962A'},
    right  ={'val':'single','sz':'18','color':'C9962A'})
cell.width = Inches(6.27)

def cover_line(text, size, bold=True, color=WHITE, space_before=0, space_after=4):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = 'Calibri'
    r.font.color.rgb = color
    return p

cover_line('', 6, space_before=18, space_after=2)
cover_line('BHS INTERNAL AUDIT', 22, color=GOLD, space_before=0, space_after=6)
cover_line('Audit Report Interface Application', 16, color=WHITE, space_before=0, space_after=4)
cover_line('Feature Reference Document', 12, bold=False, color=RGBColor(0xaa,0xc4,0xe0), space_before=0, space_after=2)
cover_line('', 6, space_before=0, space_after=14)
cover_line('Version 1.0  |  June 2025', 10, bold=False, color=RGBColor(0x88,0xaa,0xcc), space_before=0, space_after=16)

doc.add_paragraph()  # spacer

# ══════════════════════════════════════════════════════════════
#  SECTION HEADING helper
# ══════════════════════════════════════════════════════════════
section_num = [0]
def add_section(title, icon=''):
    section_num[0] += 1
    doc.add_paragraph()
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    c = tbl.cell(0, 0)
    set_cell_bg(c, NAVY)
    set_cell_border(c,
        left={'val':'single','sz':'18','color':'C9962A'})
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    r = p.add_run('{} {}  {}'.format(section_num[0], icon, title))
    r.bold = True
    r.font.size = Pt(13)
    r.font.name = 'Calibri'
    r.font.color.rgb = WHITE
    doc.add_paragraph()

def add_subsection(title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(11)
    r.font.name = 'Calibri'
    r.font.color.rgb = NAVY
    # gold underline bar
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after  = Pt(5)
    r2 = p2.add_run('─' * 60)
    r2.font.size = Pt(7)
    r2.font.color.rgb = GOLD

def add_bullet(text, sub=False, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent  = Inches(0.3 if not sub else 0.55)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(2)
    if bold_prefix:
        r1 = p.add_run(bold_prefix + ': ')
        r1.bold = True
        r1.font.size = Pt(10.5)
        r1.font.name = 'Calibri'
        r1.font.color.rgb = NAVY
    r2 = p.add_run(text)
    r2.font.size = Pt(10.5)
    r2.font.name = 'Calibri'
    r2.font.color.rgb = TEXT
    return p

def add_body(text, space_before=3, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.name = 'Calibri'
    r.font.color.rgb = TEXT
    return p

def add_note(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(6)
    p.paragraph_format.left_indent  = Inches(0.2)
    paragraph_shade(p, RGBColor(0xF0, 0xF5, 0xFF))
    r = p.add_run('ℹ  ' + text)
    r.italic = True
    r.font.size = Pt(10)
    r.font.name = 'Calibri'
    r.font.color.rgb = NAVY
    return p

# ══════════════════════════════════════════════════════════════
#  1  OVERVIEW
# ══════════════════════════════════════════════════════════════
add_section('Application Overview', '📋')
add_body(
    'The BHS Audit Report Interface is a fully self-contained, browser-based application '
    'designed to streamline the preparation, management, and communication of internal audit '
    'reports for Beaconhouse School System (BHS) branches. It requires no server, no '
    'installation, and no internet connection — open the HTML file in any modern browser '
    'and start working immediately.'
)
add_body(
    'The application integrates the complete audit workflow: building structured reports from '
    'a databank of 350+ observations, tracking work progress, assessing branch risk, generating '
    'auditee response forms, and producing final documents — all within a single file.'
)

# ── Feature summary table ──────────────────────────────────────
headers = ['Module', 'Purpose', 'Key Output']
rows_data = [
    ['Report Builder',       'Select and configure audit observations',       'Formatted audit report'],
    ['Auditee Copy',         'Generate response form for branch management',   'Printable/email-ready form'],
    ['Audit Work Status',    'Track completion progress category by category', 'Progress dashboard & charts'],
    ['Risk Assessment',      'Analyse previous reports for next-audit risk',   'Risk matrix & recommendations'],
    ['Draft by Category',    'Generate working draft reports per category',    'Category-wise draft document'],
    ['Compare Report',       'Compare current report against a saved session', 'Side-by-side gap analysis'],
    ['Response Loader',      'Import BH/RO responses into the live report',    'Merged response view'],
    ['Version History',      'Save and restore multiple report versions',      'Timestamped save slots'],
]

tbl = doc.add_table(rows=1+len(rows_data), cols=3)
tbl.style = 'Table Grid'
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
col_widths = [Inches(1.6), Inches(2.9), Inches(1.8)]
for i, w in enumerate(col_widths):
    for row in tbl.rows:
        row.cells[i].width = w

# Header row
for i, h in enumerate(headers):
    c = tbl.cell(0, i)
    set_cell_bg(c, NAVY)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    r = p.add_run(h)
    r.bold = True; r.font.size = Pt(10); r.font.name = 'Calibri'; r.font.color.rgb = WHITE

for ri, rd in enumerate(rows_data):
    row = tbl.rows[ri+1]
    bg = RGBColor(0xf3,0xf5,0xf8) if ri % 2 == 0 else WHITE
    for ci, val in enumerate(rd):
        c = row.cells[ci]
        set_cell_bg(c, bg)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after  = Pt(3)
        r = p.add_run(val)
        r.font.size = Pt(10)
        r.font.name = 'Calibri'
        r.bold = (ci == 0)
        r.font.color.rgb = NAVY if ci == 0 else TEXT

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════
#  2  REPORT BUILDER
# ══════════════════════════════════════════════════════════════
add_section('Report Builder', '📝')
add_body(
    'The Report Builder is the core of the application. Auditors use it to construct a '
    'complete, professionally formatted audit report by selecting relevant observations from '
    'the BHS databank and configuring branch-specific information.'
)

add_subsection('Branch & Report Information')
add_bullet('Branch code, name, region, and city with auto-fill from the built-in branch directory')
add_bullet('Audit type (Routine, Special, Follow-up, etc.)')
add_bullet('Report date, audit period (from/to), visit dates, and previous audit date')
add_bullet('Audit framework reference (e.g. COSO Internal Control Framework)')
add_bullet('Prepared By — configurable auditor name dropdown')
add_note('All branch fields are auto-locked once the report is generated to prevent accidental changes. An override unlock button is available.')

add_subsection('Observation Databank (350+ Entries)')
add_bullet('Over 350 pre-defined observations across 9 audit categories')
add_bullet('Each observation includes: title, detailed finding, criteria/requirement, recommendation, and department/process tags')
add_bullet('Full-text search with keyword scoring — finds relevant observations as you type')
add_bullet('Category filter to browse specific audit areas')
add_bullet('Type classification: New or Repeated — displayed as colour-coded pills on each selected observation')
add_note('The nine canonical categories are: Branch Imprest/Payments & Accounts; Cash & Bank; Recruitment, Payroll & Employee Record; Student Record; Fixed Assets; Library; Science Laboratory; Safety & Security; Others.')

add_subsection('Observation Configuration Per Entry')
add_bullet('Inline title and body editing — modify any observation text for branch-specific accuracy')
add_bullet('Sensitivity level: Routine / Significant / Critical')
add_bullet('COSO component tagging (Control Environment, Risk Assessment, Control Activities, etc.)')
add_bullet('Auditee Management Action input (committed resolution step)')
add_bullet('Added Value / Benefit statement field')
add_bullet('Reference document attachment field')
add_bullet('Annexure linking — attach supporting evidence annexures to any observation')

add_subsection('Report Generation')
add_bullet('Pending-change queue — additions and removals are staged before confirmation, preventing accidental report modifications')
add_bullet('One-click Generate / Refresh Report renders all selected observations into a formatted report view')
add_bullet('Report automatically groups observations by category in the canonical order')
add_bullet('Each category group shows: category heading, sub-heading, observation count, repeated/new split')
add_bullet('Two classification sections per report: New Observations and Repeated Observations')
add_bullet('Annexure section rendered at the end with all linked evidence documents')

add_subsection('Report Export Options')
add_bullet('Copy to Clipboard — copies the full formatted report as rich text for pasting into Word or email')
add_bullet('Print / Save as PDF — the print layout hides all UI chrome and formats the report cleanly for A4 paper')
add_bullet('Finalize & Save — stores the completed report with a timestamp and branch metadata')

# ══════════════════════════════════════════════════════════════
#  3  AUDITEE COPY
# ══════════════════════════════════════════════════════════════
add_section('Auditee Response Form', '📧')
add_body(
    'The Auditee Copy feature generates a standalone response document that is sent to branch '
    'management after the audit. It presents each observation in a readable format and provides '
    'structured fields for the branch to formally respond.'
)
add_bullet('Generates a self-contained HTML file that can be printed, emailed, or filled in a browser')
add_bullet('Each observation is displayed with its full finding, criteria, and recommendation')
add_bullet('Branch management fills in a free-text Response / Comments box per observation')
add_bullet('Estimated Time to Resolve field — branch commits to a specific timeframe (e.g. "2 weeks", "by 31-Dec-2025")')
add_bullet('A single Download Response button packages all responses into a JSON file')
add_bullet('Two variants: Branch Head (BH) copy and Regional Office/Head Office (RO/HO) copy — each loads separately without affecting the other')
add_note('When a response file is loaded back into the Report Builder, the responses appear under the correct BH or RO heading only — cross-contamination between the two response streams has been fully prevented.')

# ══════════════════════════════════════════════════════════════
#  4  AUDIT WORK STATUS
# ══════════════════════════════════════════════════════════════
add_section('Audit Work Status', '✅')
add_body(
    'The Audit Work Status panel is an embedded checklist tool that lets the audit team '
    'monitor their own field work progress in real time — tracking which observations have '
    'been checked, which are clear, and which carry findings.'
)

add_subsection('Category-wise Observation Checklist')
add_bullet('All 350+ observations listed category by category in the canonical order')
add_bullet('Each observation has a status dropdown with five states:')
add_bullet('Not Started — work has not begun on this item', sub=True)
add_bullet('In Progress — currently being examined', sub=True)
add_bullet('Checked with Observation — reviewed; a finding has been identified (amber highlight)', sub=True)
add_bullet('Checked & Clear — reviewed; no issue found (green highlight)', sub=True)
add_bullet('Not Applicable — this observation does not apply to this branch', sub=True)
add_bullet('Remark field per observation for internal audit team notes')
add_bullet('Auto-sync with Report Builder: observations already added to the audit report are automatically marked "Checked with Observation" when the Work Status panel opens')

add_subsection('Progress Dashboard')
add_bullet('Heat-map grid: one tile per category, colour-coded from red (0%) to green (100%)')
add_bullet('Each tile shows the completion percentage and observation count for that category')
add_bullet('Click any heat-map tile to scroll directly to that category in the checklist')
add_bullet('Overall progress bar showing total checked vs. total observations')

add_subsection('Progress Report')
add_bullet('Tabular report: category-wise counts of each status type with colour-coded cells')
add_bullet('Visual bar chart showing percentage done per category')
add_bullet('Branch info bar at the top showing branch code, name, region, and city pulled automatically from the main report fields')
add_note('Work Status data is stored in browser localStorage (key: bhs_aws_v1_data) and persists between sessions on the same browser.')

# ══════════════════════════════════════════════════════════════
#  5  RISK ASSESSMENT
# ══════════════════════════════════════════════════════════════
add_section('Risk Assessment Tool', '📊')
add_body(
    'The Risk Assessment tab analyses previous audit reports to compute a risk profile for '
    'each category and generate prioritised focus recommendations for the upcoming audit visit.'
)

add_subsection('Data Input')
add_bullet('Branch selector with auto-fill of region and city from the built-in branch directory')
add_bullet('File loader supporting Word (.docx), Excel (.xlsx), JSON session exports, and plain text')
add_bullet('Multiple files can be loaded simultaneously for trend analysis across two audit cycles')
add_bullet('Paste text directly into a text area if no digital file is available')
add_bullet('Audit type and Prepared By fields for report header customisation')

add_subsection('Risk Scoring Engine')
add_bullet('Seven risk categories analysed: Financial Management, HR & Payroll, Safety & Security, Fixed Assets & Inventory, Library & Records, Student Records & Fee, General Administration')
add_bullet('Each category is scored 0–100 using four weighted factors:')
add_bullet('Observation Frequency (40 pts): how many findings landed in this category relative to total', sub=True)
add_bullet('Recurrence Score (30 pts): proportion of findings that are repeated from prior audits', sub=True)
add_bullet('Databank Weight Score (20 pts): how heavily this category features in the BHS historical databank', sub=True)
add_bullet('Severity Bonus (10 pts): Safety & Security receives an automatic inherent-risk premium', sub=True)
add_bullet('Score bands: Low (0–24), Medium (25–49), High (50–74), Critical (75–100)')

add_subsection('Risk Output Panels')
add_bullet('Overall Risk Score gauge — needle indicator on a 0–100 colour-gradient scale')
add_bullet('Summary tiles: count of categories at each risk level (Critical / High / Medium / Low)')
add_bullet('Category Risk Matrix table with score bar, risk badge, observation count, and databank count')
add_bullet('Observation Frequency Bar Chart — horizontal bars per category with repeated-observation callout')
add_bullet('Priority Focus Areas cards (top 5 categories) with explanation, score, and BHS databank checkpoints')
add_bullet('Trend Analysis (when two reports are loaded): category-by-category worsening/improving arrows')
add_bullet('Persistent / Repeated Issues card highlighting unresolved findings from prior audits')
add_bullet('Specific Audit Recommendations numbered list with priority labelling and targeted checkpoints')

# ══════════════════════════════════════════════════════════════
#  6  ADDITIONAL MODULES
# ══════════════════════════════════════════════════════════════
add_section('Additional Modules', '🔧')

add_subsection('Draft by Category')
add_body('Generates a printable working draft report for a single selected category, useful for field reviews and category-level team assignments.')
add_bullet('Category selector — choose any of the nine audit categories')
add_bullet('Renders all selected observations for that category in a clean, print-ready layout')
add_bullet('Includes branch info header and observation details with blank response fields for manual notes')

add_subsection('Compare Report')
add_body('Allows the auditor to load a previously saved JSON session and compare it side-by-side against the current report.')
add_bullet('Drag-and-drop or file-browse to load a saved session JSON')
add_bullet('Category and status filters to focus the comparison')
add_bullet('Colour-coded results: green = observation present in both; red = present in old but missing from new')
add_bullet('Count badges per category showing how many observations match or are missing')

add_subsection('Load Response')
add_body('Imports a JSON response file submitted by the branch (BH or RO/HO) into the live report.')
add_bullet('Select the response type (Branch Head or Regional/Head Office) before loading')
add_bullet('Responses populate automatically under the correct heading in the rendered report')
add_bullet('Loading one type does not affect the other — both can coexist in the same report')

add_subsection('Load Report')
add_body('Restores a previously exported report session, allowing work to continue across sessions or devices.')
add_bullet('Accepts the JSON session file exported by Finalize & Save or the version history')
add_bullet('Fully restores: selected observations, branch info, type overrides, custom observations, responses, and annexures')

add_subsection('Version History')
add_body('Manages multiple saved snapshots of the report within browser localStorage.')
add_bullet('Up to several versions stored with timestamps and branch name')
add_bullet('One-click restore to any previous version')
add_bullet('Export any saved version as a standalone JSON session file')

# ══════════════════════════════════════════════════════════════
#  7  PRESENTATION FEATURES
# ══════════════════════════════════════════════════════════════
add_section('Presentation & Usability Features', '✨')

add_subsection('Animated Splash Screen')
add_body('A branded full-screen intro overlay appears on application load.')
add_bullet('BHS Internal Audit logo with animated pulse effect')
add_bullet('Animated gold progress bar that fills over 2.8 seconds')
add_bullet('Auto-dismisses after 2.8 seconds; click anywhere to skip immediately')

add_subsection('Demo Mode')
add_body('A single-click Demo button in the toolbar loads a complete realistic sample report for demonstration purposes.')
add_bullet('Pre-fills: LMA 11-FCC Gulberg, Lahore (Branch 240, Center Region)')
add_bullet('14 observations selected across 7 categories — a mix of New and Repeated findings')
add_bullet('Renders the full report instantly with no manual data entry')
add_bullet('Ideal for training sessions, stakeholder presentations, and system demonstrations')

add_subsection('Visual Design')
add_bullet('Navy-to-dark-navy gradient toolbar with a gold accent border')
add_bullet('Smooth hover animations on all toolbar buttons (lift + shadow effect)')
add_bullet('Panel fade-in animation when switching between Report, Work Status, and Risk Assessment views')
add_bullet('Gold focus rings on all form inputs and dropdowns for accessibility')
add_bullet('Colour-coded observation pills: green for New observations, red/amber for Repeated')
add_bullet('Responsive layout adapts from wide-screen workstations to laptop displays')

add_subsection('Print / PDF Export')
add_bullet('Dedicated @media print CSS hides all UI chrome (toolbar, left panel, controls)')
add_bullet('A4 page setup with 15 mm margins via @page rule')
add_bullet('Proper page breaks around observation groups to prevent awkward splits')
add_bullet('Category header colours preserved in PDF via print-color-adjust property')
add_bullet('Clean observation blocks with border-left type indicator lines')

# ══════════════════════════════════════════════════════════════
#  8  TECHNICAL SPECIFICATIONS
# ══════════════════════════════════════════════════════════════
add_section('Technical Specifications', '⚙')

tbl2 = doc.add_table(rows=9, cols=2)
tbl2.style = 'Table Grid'
tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
specs = [
    ('File Format',         'Single self-contained HTML file — no dependencies, no CDN, no server required'),
    ('Browser Support',     'Chrome 90+, Edge 90+, Firefox 88+, Safari 14+ (any modern browser)'),
    ('JavaScript Standard', 'ES5-compatible — maximum cross-browser compatibility'),
    ('Data Storage',        'Browser localStorage for Work Status and report versions; JSON file export/import for portability'),
    ('Observation Databank','350+ pre-coded observations across 9 categories built into the application'),
    ('Branch Directory',    '200+ BHS branches with code, name, region, city, and head-of-campus data'),
    ('File Size',           'Approximately 500 KB — loads instantly even without internet'),
    ('Offline Operation',   'Fully offline — no data is ever sent to any server or external service'),
    ('Print Format',        'A4 optimised print layout; compatible with browser Print-to-PDF'),
]
for ri, (k, v) in enumerate(specs):
    row = tbl2.rows[ri]
    bg = RGBColor(0xf3,0xf5,0xf8) if ri % 2 == 0 else WHITE
    set_cell_bg(row.cells[0], NAVY if ri == 0 else bg)
    set_cell_bg(row.cells[1], NAVY if ri == 0 else bg)
    for ci, val in enumerate([k, v]):
        c = row.cells[ci]
        if ri > 0: set_cell_bg(c, bg)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after  = Pt(4)
        r = p.add_run(val)
        r.font.size = Pt(10)
        r.font.name = 'Calibri'
        r.bold = (ci == 0)
        r.font.color.rgb = NAVY if ci == 0 else TEXT

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════
#  FOOTER NOTE
# ══════════════════════════════════════════════════════════════
p_footer = doc.add_paragraph()
p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_footer.paragraph_format.space_before = Pt(16)
r_footer = p_footer.add_run('BHS Internal Audit — Audit Report Interface Application  |  Feature Reference Document  |  June 2025')
r_footer.italic = True
r_footer.font.size = Pt(9)
r_footer.font.name = 'Calibri'
r_footer.font.color.rgb = SUBTEXT

# ── Save ──────────────────────────────────────────────────────
out = '/home/user/KS324/BHS_Audit_Report_Interface_Features.docx'
doc.save(out)
print('Saved:', out)
