from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page setup ────────────────────────────────────────────────
sec = doc.sections[0]
sec.page_width    = Inches(8.27)
sec.page_height   = Inches(11.69)
sec.left_margin   = Inches(1.0)
sec.right_margin  = Inches(1.0)
sec.top_margin    = Inches(0.85)
sec.bottom_margin = Inches(0.85)

# ── Colours ───────────────────────────────────────────────────
NAVY    = RGBColor(0x1a, 0x3a, 0x5c)
DARK    = RGBColor(0x0f, 0x25, 0x40)
GOLD    = RGBColor(0xc9, 0x96, 0x2a)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
TEXT    = RGBColor(0x1e, 0x2d, 0x40)
MUTED   = RGBColor(0x6a, 0x78, 0x88)
GREEN   = RGBColor(0x1a, 0x60, 0x30)
GREENBG = RGBColor(0xe8, 0xf5, 0xec)
GOLDBG  = RGBColor(0xff, 0xf8, 0xe8)
NAVYBG  = RGBColor(0xf0, 0xf4, 0xfa)

def hex_str(rgb): return '{:02X}{:02X}{:02X}'.format(rgb[0], rgb[1], rgb[2])

def set_cell_bg(cell, rgb):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_str(rgb))
    tcPr.append(shd)

def set_borders(cell, sides):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcB = OxmlElement('w:tcBorders')
    for side, cfg in sides.items():
        el = OxmlElement('w:' + side)
        el.set(qn('w:val'),   cfg.get('val','single'))
        el.set(qn('w:sz'),    cfg.get('sz','6'))
        el.set(qn('w:space'),'0')
        el.set(qn('w:color'), cfg.get('color','auto'))
        tcB.append(el)
    tcPr.append(tcB)

def no_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcB = OxmlElement('w:tcBorders')
    for side in ['top','bottom','left','right','insideH','insideV']:
        el = OxmlElement('w:' + side)
        el.set(qn('w:val'), 'none')
        el.set(qn('w:sz'),  '0')
        el.set(qn('w:space'),'0')
        el.set(qn('w:color'),'auto')
        tcB.append(el)
    tcPr.append(tcB)

def run(para, text, bold=False, italic=False, size=11, color=TEXT, font='Calibri'):
    r = para.add_run(text)
    r.bold = bold; r.italic = italic
    r.font.size = Pt(size); r.font.name = font
    r.font.color.rgb = color
    return r

def para(doc, text='', bold=False, italic=False, size=11, color=TEXT,
         align=WD_ALIGN_PARAGRAPH.LEFT, sb=0, sa=6, font='Calibri'):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after  = Pt(sa)
    if text:
        r = p.add_run(text)
        r.bold = bold; r.italic = italic
        r.font.size = Pt(size); r.font.name = font
        r.font.color.rgb = color
    return p

# ══════════════════════════════════════════════════════════════
#  COVER BLOCK
# ══════════════════════════════════════════════════════════════
cover = doc.add_table(rows=1, cols=1)
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = cover.cell(0, 0)
set_cell_bg(cell, DARK)
set_borders(cell, {
    'top':    {'val':'single','sz':'24','color':'C9962A'},
    'bottom': {'val':'single','sz':'24','color':'C9962A'},
    'left':   {'val':'single','sz':'24','color':'C9962A'},
    'right':  {'val':'single','sz':'24','color':'C9962A'},
})

def cline(text, size, bold=True, color=WHITE, sb=0, sa=4):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after  = Pt(sa)
    r = p.add_run(text)
    r.bold = bold; r.font.size = Pt(size)
    r.font.name = 'Calibri'; r.font.color.rgb = color

cline('', 4, sb=20, sa=2)
cline('BEACONHOUSE SCHOOL SYSTEM', 11, bold=False,
      color=RGBColor(0x88,0xaa,0xcc), sb=0, sa=3)
cline('BHS Internal Audit', 24, color=GOLD, sb=0, sa=5)
cline('Audit Report Interface Application', 16, sb=0, sa=4)
cline('Executive Summary for Senior Management', 12, bold=False,
      color=RGBColor(0xaa,0xc4,0xe0), sb=0, sa=3)

# gold divider
pd = cell.add_paragraph()
pd.alignment = WD_ALIGN_PARAGRAPH.CENTER
pd.paragraph_format.space_before = Pt(6)
pd.paragraph_format.space_after  = Pt(6)
r = pd.add_run('━' * 38)
r.font.color.rgb = GOLD; r.font.size = Pt(9); r.font.name = 'Calibri'

cline('Prepared by: BHS Internal Audit Department', 10, bold=False,
      color=RGBColor(0x88,0xaa,0xcc), sb=0, sa=2)
cline('June 2025', 10, bold=False,
      color=RGBColor(0x88,0xaa,0xcc), sb=0, sa=18)

doc.add_paragraph()   # spacer

# ══════════════════════════════════════════════════════════════
#  PURPOSE STATEMENT — full-width tinted box
# ══════════════════════════════════════════════════════════════
pt = doc.add_table(rows=1, cols=1)
pt.alignment = WD_TABLE_ALIGNMENT.CENTER
pc = pt.cell(0, 0)
set_cell_bg(pc, NAVYBG)
set_borders(pc, {'left': {'val':'single','sz':'18','color':'1A3A5C'}})
no_borders(pc)

pp = pc.add_paragraph()
pp.paragraph_format.space_before = Pt(8)
pp.paragraph_format.space_after  = Pt(8)
pp.paragraph_format.left_indent  = Inches(0.1)
run(pp, 'Purpose:  ', bold=True, size=11, color=NAVY)
run(pp,
    'This document provides senior management with a concise overview of the BHS Audit Report '
    'Interface — a purpose-built digital tool that modernises and standardises the internal audit '
    'process across all BHS branches.',
    size=11, color=TEXT)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════
#  SECTION HEADING helper
# ══════════════════════════════════════════════════════════════
sn = [0]
def heading(title, icon=''):
    sn[0] += 1
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    set_cell_bg(c, NAVY)
    no_borders(c)
    set_borders(c, {'left': {'val':'single','sz':'24','color':'C9962A'}})
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after  = Pt(5)
    r = p.add_run('  {}  {}  {}'.format(sn[0], icon, title))
    r.bold = True; r.font.size = Pt(12)
    r.font.name = 'Calibri'; r.font.color.rgb = WHITE
    doc.add_paragraph()

def kv_bullet(label, value, label_color=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Inches(0.25)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(3)
    # bullet mark
    rb = p.add_run('◆  ')
    rb.font.size = Pt(7); rb.font.color.rgb = GOLD; rb.font.name = 'Calibri'
    run(p, label + ':  ', bold=True, size=10.5, color=label_color)
    run(p, value, size=10.5, color=TEXT)

def body(text, sb=3, sa=5, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after  = Pt(sa)
    if indent: p.paragraph_format.left_indent = Inches(indent)
    r = p.add_run(text)
    r.font.size = Pt(10.5); r.font.name = 'Calibri'; r.font.color.rgb = TEXT
    return p

# ══════════════════════════════════════════════════════════════
#  1  BACKGROUND
# ══════════════════════════════════════════════════════════════
heading('Background', '📌')
body(
    'The BHS Internal Audit function is responsible for reviewing financial controls, staff records, '
    'student data, fixed assets, library management, security arrangements, and general administration '
    'across hundreds of BHS branches nationwide. Historically, audit reports were prepared manually '
    'using word-processing templates, a process that was time-consuming, prone to inconsistency, and '
    'offered limited visibility into audit progress or risk trends.'
)
body(
    'To address these challenges, the BHS Audit Report Interface Application was developed — a single, '
    'self-contained digital tool that covers the entire audit lifecycle from fieldwork to final report.'
)

# ══════════════════════════════════════════════════════════════
#  2  WHAT THE APPLICATION DOES
# ══════════════════════════════════════════════════════════════
heading('What the Application Does', '🔍')
body(
    'The application is a single HTML file that opens in any web browser — no installation, no internet '
    'connection, and no server are required. It brings together six core functions:'
)

# 2x3 feature cards table
cards = [
    ('📝', 'Report Building',      'Auditors select from 350+ pre-coded observations and generate a fully formatted audit report within minutes.'),
    ('✅', 'Work Progress Tracking','A live checklist tracks which observations have been examined, producing category-wise completion dashboards and charts.'),
    ('📊', 'Risk Assessment',       'Previous audit reports are analysed to score risk per category and generate prioritised focus areas for the next visit.'),
    ('📧', 'Auditee Communication', 'A formal response form is generated for Branch Head and Regional Office/Head Office, capturing commitments and resolution timelines.'),
    ('📂', 'Document Management',   'Reports are saved as versioned snapshots, compared against prior audits, and exported as JSON or PDF.'),
    ('🖨', 'Report Export',         'Completed reports are exported as polished PDFs or copied as rich text for pasting into official correspondence.'),
]
tbl = doc.add_table(rows=3, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for ri in range(3):
    for ci in range(2):
        idx = ri*2 + ci
        icon, title, desc = cards[idx]
        c = tbl.cell(ri, ci)
        set_cell_bg(c, NAVYBG if (ri+ci)%2==0 else GOLDBG)
        no_borders(c)
        c.width = Inches(3.1)
        p1 = c.add_paragraph()
        p1.paragraph_format.space_before = Pt(8)
        p1.paragraph_format.space_after  = Pt(2)
        p1.paragraph_format.left_indent  = Inches(0.1)
        run(p1, icon + '  ' + title, bold=True, size=11, color=NAVY)
        p2 = c.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after  = Pt(8)
        p2.paragraph_format.left_indent  = Inches(0.1)
        run(p2, desc, size=10, color=TEXT)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════
#  3  KEY BENEFITS FOR MANAGEMENT
# ══════════════════════════════════════════════════════════════
heading('Key Benefits for Management', '📈')

benefits = [
    ('Standardised Reporting',
     'Every report follows the same structure, category order, and language — eliminating inconsistencies across auditors and audit cycles.'),
    ('Faster Turnaround',
     'Report preparation time is significantly reduced. Selecting and configuring observations takes minutes rather than hours of manual drafting.'),
    ('Real-Time Progress Visibility',
     'The Audit Work Status module gives audit managers an instant view of fieldwork completion — by category, by status, and as an overall percentage.'),
    ('Structured Accountability',
     'Auditee response forms capture formal management commitments, including the estimated timeframe to resolve each observation, creating a clear audit trail.'),
    ('Risk-Informed Planning',
     'The Risk Assessment tool analyses historical findings to score each category\'s risk level (Low / Medium / High / Critical) and recommend priority focus areas for the next visit.'),
    ('Trend Analysis',
     'Repeated observations are automatically flagged and tracked, allowing management to identify persistent control weaknesses across multiple audit cycles.'),
    ('Full Audit Trail',
     'Every version of a report is stored with a timestamp. Reports can be exported, shared, and reloaded — ensuring nothing is lost between sessions.'),
    ('Zero IT Infrastructure Cost',
     'The tool requires no server, no database, no installation, and no subscription. It runs entirely in the auditor\'s browser from a single shared file.'),
]

for label, value in benefits:
    kv_bullet(label, value)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════
#  4  COVERAGE & SCALE
# ══════════════════════════════════════════════════════════════
heading('Coverage & Scale', '🗺')

# Stats table
stats = [
    ('350+',  'Pre-coded audit\nobservations'),
    ('9',     'Audit categories\ncovered'),
    ('200+',  'BHS branches in\nbuilt-in directory'),
    ('5',     'Work-status states\nper observation'),
]
st = doc.add_table(rows=1, cols=4)
st.alignment = WD_TABLE_ALIGNMENT.CENTER
for ci, (num, lbl) in enumerate(stats):
    c = st.cell(0, ci)
    c.width = Inches(1.5)
    set_cell_bg(c, NAVY)
    set_borders(c, {
        'left':  {'val':'single','sz':'2','color':'C9962A'},
        'right': {'val':'single','sz':'2','color':'C9962A'},
    })
    p1 = c.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(10)
    p1.paragraph_format.space_after  = Pt(2)
    run(p1, num, bold=True, size=22, color=GOLD)
    p2 = c.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after  = Pt(10)
    run(p2, lbl, size=9, color=RGBColor(0xaa,0xc4,0xe0))

doc.add_paragraph()

body('The nine audit categories covered are: Branch Imprest, Payments & Accounts; Cash & Bank; '
     'Recruitment, Payroll & Employee Record; Students\' Record; Fixed Assets; Library; '
     'Science Laboratory; Safety & Security; and Others. Category ordering in all reports '
     'follows this canonical sequence consistently.')

# ══════════════════════════════════════════════════════════════
#  5  AUDIT WORKFLOW
# ══════════════════════════════════════════════════════════════
heading('Supported Audit Workflow', '🔄')

steps = [
    ('Pre-Audit',   'Risk Assessment',
     'Auditor loads previous audit reports or sessions. The system scores each category '
     'by risk level and recommends priority areas. This informs the audit plan before fieldwork begins.'),
    ('Fieldwork',   'Work Status Tracking',
     'During the visit, auditors mark each observation as Not Started, In Progress, '
     'Checked & Clear, Checked with Observation, or Not Applicable. Progress dashboards '
     'update in real time.'),
    ('Reporting',   'Report Builder',
     'Selected observations are configured and generated into a fully formatted report. '
     'Repeated findings are highlighted. The report is reviewed, refined, and finalised.'),
    ('Issuance',    'Auditee Copy',
     'A structured response form is sent to the Branch Head and Regional/Head Office. '
     'Management responses and resolution timelines are captured.'),
    ('Follow-up',   'Response Loading',
     'Submitted responses are loaded back into the live report. The final report — with '
     'management comments — is saved, exported as PDF, and archived.'),
]

for phase, module, desc in steps:
    wt = doc.add_table(rows=1, cols=3)
    wt.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(wt.cell(0,0)); no_borders(wt.cell(0,1)); no_borders(wt.cell(0,2))
    # Phase badge
    c0 = wt.cell(0, 0); c0.width = Inches(1.0)
    set_cell_bg(c0, NAVY)
    p0 = c0.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(6)
    p0.paragraph_format.space_after  = Pt(6)
    run(p0, phase, bold=True, size=10, color=GOLD)
    # Module
    c1 = wt.cell(0, 1); c1.width = Inches(1.4)
    set_cell_bg(c1, NAVYBG)
    p1 = c1.add_paragraph()
    p1.paragraph_format.space_before = Pt(6)
    p1.paragraph_format.space_after  = Pt(6)
    p1.paragraph_format.left_indent  = Inches(0.1)
    run(p1, module, bold=True, size=10, color=NAVY)
    # Description
    c2 = wt.cell(0, 2); c2.width = Inches(3.87)
    p2 = c2.add_paragraph()
    p2.paragraph_format.space_before = Pt(6)
    p2.paragraph_format.space_after  = Pt(6)
    p2.paragraph_format.left_indent  = Inches(0.1)
    run(p2, desc, size=10, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ══════════════════════════════════════════════════════════════
#  6  CONCLUSION & RECOMMENDATION
# ══════════════════════════════════════════════════════════════
heading('Conclusion & Management Recommendation', '✅')

# Tinted conclusion box
ct = doc.add_table(rows=1, cols=1)
ct.alignment = WD_TABLE_ALIGNMENT.CENTER
cc = ct.cell(0, 0)
set_cell_bg(cc, GREENBG)
no_borders(cc)
set_borders(cc, {'left': {'val':'single','sz':'18','color':'1A6030'}})
cp = cc.add_paragraph()
cp.paragraph_format.space_before = Pt(8)
cp.paragraph_format.space_after  = Pt(8)
cp.paragraph_format.left_indent  = Inches(0.1)
run(cp,
    'The BHS Audit Report Interface Application represents a significant step forward in the '
    'professionalism, efficiency, and consistency of BHS internal audit operations. By digitising '
    'the full audit lifecycle — from risk profiling and fieldwork tracking to report generation '
    'and auditee response management — the tool enables the audit function to deliver higher-quality '
    'outputs with greater speed and less administrative effort.',
    size=10.5, color=GREEN)

doc.add_paragraph()
body('Management is recommended to:')
for item in [
    'Formally adopt the application as the standard tool for all BHS internal audit assignments.',
    'Ensure all audit team members receive orientation on its features and workflow.',
    'Require branch management to submit responses through the structured Auditee Copy form to enable consistent response tracking.',
    'Utilise the Risk Assessment output when setting audit priorities and allocating audit resources across regions.',
    'Review the Audit Work Status progress dashboard during audit visits to monitor real-time fieldwork completion.',
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Inches(0.25)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(3)
    rb = p.add_run('✔  ')
    rb.font.size = Pt(10); rb.font.color.rgb = GREEN; rb.font.name = 'Calibri'
    run(p, item, size=10.5, color=TEXT)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════
#  FOOTER
# ══════════════════════════════════════════════════════════════
ft = doc.add_table(rows=1, cols=1)
ft.alignment = WD_TABLE_ALIGNMENT.CENTER
fc = ft.cell(0, 0)
set_cell_bg(fc, DARK)
no_borders(fc)
fp = fc.add_paragraph()
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.space_before = Pt(7)
fp.paragraph_format.space_after  = Pt(7)
run(fp, 'BHS Internal Audit Department  |  Audit Report Interface — Executive Summary  |  June 2025',
    size=9, bold=False, color=RGBColor(0x88,0xaa,0xcc))

# ── Save ──────────────────────────────────────────────────────
out = '/home/user/KS324/BHS_Audit_Application_Executive_Summary.docx'
doc.save(out)
print('Saved:', out)
